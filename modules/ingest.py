"""Turn an uploaded file into plain text. Returns (text, how_it_was_read)."""
import io
from pathlib import Path

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
TEXT_EXT = {".txt", ".md", ".csv"}
ALLOWED = ["pdf", "docx", "txt", "md", "csv", "png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff"]


def _ocr_image(img):
    """Tesseract first (free, local). Falls back to Claude vision if a key exists."""
    try:
        import pytesseract
        for lang in ("eng+hin+mar", "eng"):
            try:
                return pytesseract.image_to_string(img, lang=lang), "Read with OCR (Tesseract)"
            except Exception:
                continue
    except Exception:
        pass
    from modules import llm
    if llm.available():
        buf = io.BytesIO()
        img.convert("RGB").save(buf, format="PNG")
        txt = llm.vision_text(buf.getvalue())
        if txt:
            return txt, "Read with AI vision"
    raise RuntimeError("No OCR engine available. Please paste the report text instead.")


def read_upload(uploaded):
    name = uploaded.name
    ext = Path(name).suffix.lower()
    data = uploaded.getvalue()

    if ext in TEXT_EXT:
        return data.decode("utf-8", errors="ignore"), "Text file"

    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append("  ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines), "Word document"

    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = pdf.pages[:15]
            text = "\n".join((p.extract_text() or "") for p in pages)
            if len(text.strip()) >= 30:
                return text, "PDF text"
            parts = [_ocr_image(p.to_image(resolution=200).original)[0] for p in pages[:3]]
            return "\n".join(parts), "Scanned PDF read with OCR (first 3 pages)"

    if ext in IMAGE_EXT:
        from PIL import Image
        return _ocr_image(Image.open(io.BytesIO(data)))

    raise ValueError(f"Unsupported file type: {ext}")
